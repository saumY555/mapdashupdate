"""
report_exporter.py
Exports an analysis report to PDF, Word (.docx), or CSV formats.
"""

import io
import csv
import re
from typing import Dict


# ── PDF Export ────────────────────────────────────────────────────────────────

def export_to_pdf(report: Dict) -> bytes:
    """Export the analysis report as a styled PDF."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_fill_color(30, 30, 60)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 14, "Document Analysis Report", new_x="LMARGIN", new_y="NEXT", fill=True, align="C")
    pdf.ln(4)

    # Filename & stats
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(0, 8, f"Document: {report['filename']}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, f"Pages/Sections: {report['page_count']}  |  Characters: {report['char_count']:,}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    # Report content
    pdf.set_text_color(0, 0, 0)
    raw = report["raw_report"]

    for line in raw.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue

        # Section headers (## Heading)
        if line.startswith("## "):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 13)
            pdf.set_text_color(30, 30, 120)
            safe = line[3:].encode("latin-1", errors="replace").decode("latin-1")
            pdf.cell(0, 10, safe, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "", 11)

        # Bullet points
        elif line.startswith("- ") or line.startswith("* "):
            text = line[2:]
            # Strip bold markdown **text**
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            safe = text.encode("latin-1", errors="replace").decode("latin-1")
            pdf.cell(6)
            pdf.multi_cell(0, 7, f"\u2022  {safe}")

        # Numbered list
        elif re.match(r"^\d+\.", line):
            safe = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            safe = safe.encode("latin-1", errors="replace").decode("latin-1")
            pdf.cell(6)
            pdf.multi_cell(0, 7, safe)

        # Normal text
        else:
            safe = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            safe = safe.encode("latin-1", errors="replace").decode("latin-1")
            pdf.multi_cell(0, 7, safe)

    return pdf.output()


# ── Word Export ───────────────────────────────────────────────────────────────

def export_to_docx(report: Dict) -> bytes:
    """Export the analysis report as a Word .docx file."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading("Document Analysis Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Meta
    meta = doc.add_paragraph()
    meta.add_run(f"Document: ").bold = True
    meta.add_run(report["filename"])
    meta2 = doc.add_paragraph()
    meta2.add_run(f"Pages/Sections: {report['page_count']}  |  Characters: {report['char_count']:,}")
    doc.add_paragraph()

    raw = report["raw_report"]

    for line in raw.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)

        elif line.startswith("- ") or line.startswith("* "):
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:])
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)

        elif re.match(r"^\d+\.", line):
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\.\s*", "", text))

        else:
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── CSV Export ────────────────────────────────────────────────────────────────

def export_to_csv(report: Dict) -> bytes:
    """Export key points and entities from the report as a CSV."""
    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(["Section", "Content"])
    writer.writerow(["Document", report["filename"]])
    writer.writerow(["Pages/Sections", report["page_count"]])
    writer.writerow(["Characters", report["char_count"]])
    writer.writerow([])

    current_section = ""
    for line in report["raw_report"].split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("## "):
            current_section = line[3:]
        elif line.startswith("- ") or line.startswith("* "):
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:])
            writer.writerow([current_section, text])
        elif re.match(r"^\d+\.", line):
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            writer.writerow([current_section, text])
        else:
            writer.writerow([current_section, re.sub(r"\*\*(.*?)\*\*", r"\1", line)])

    return buf.getvalue().encode("utf-8-sig")  # utf-8-sig for Excel compatibility
